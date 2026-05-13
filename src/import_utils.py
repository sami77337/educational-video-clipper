"""Import clip table rows from Excel and CSV files."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.exclusions import format_exclusions, parse_exclusions
from src.time_utils import format_seconds, normalize_digits, normalize_timestamp_text


REQUIRED_COLUMNS = ("number", "title", "start", "end")
OPTIONAL_COLUMNS = ("exclusions",)
AR_MISSING_REQUIRED_COLUMNS = "ملف Excel لا يحتوي على الأعمدة المطلوبة: number, title, start, end"
AR_UNSUPPORTED_IMPORT_FILE = "صيغة الملف غير مدعومة. استخدم xlsx أو csv"
AR_XLS_REQUIRES_CONVERSION = "صيغة .xls غير مدعومة حاليًا. احفظ الملف بصيغة .xlsx أو .csv"
AR_INVALID_TIME_VALUE = "قيمة الوقت غير صحيحة"


@dataclass(frozen=True)
class ImportedClipRow:
    """A clip row imported from a spreadsheet-like file."""

    number: int
    title: str
    start: str
    end: str
    exclusions: str = ""


class ClipImportError(ValueError):
    """Raised when an import file cannot be read or validated."""


def import_clip_rows(file_path: str | Path) -> list[ImportedClipRow]:
    """Import clip rows from .xlsx, .xls if possible, or .csv."""

    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return import_xlsx_clip_rows(path)
    if suffix == ".csv":
        return import_csv_clip_rows(path)
    if suffix == ".xls":
        return import_xls_clip_rows(path)

    raise ClipImportError(AR_UNSUPPORTED_IMPORT_FILE)


def import_xlsx_clip_rows(file_path: str | Path) -> list[ImportedClipRow]:
    """Import rows from an .xlsx workbook."""

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as error:
            raise ClipImportError(AR_MISSING_REQUIRED_COLUMNS) from error

        column_map = _build_column_map(headers)
        return [
            _row_from_mapping(_values_to_mapping(row, column_map), index)
            for index, row in enumerate(rows, start=2)
            if _row_has_data(row)
        ]
    finally:
        workbook.close()


def import_csv_clip_rows(file_path: str | Path) -> list[ImportedClipRow]:
    """Import rows from a CSV file."""

    with Path(file_path).open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.DictReader(csv_file)
        if reader.fieldnames is None:
            raise ClipImportError(AR_MISSING_REQUIRED_COLUMNS)

        _build_column_map(reader.fieldnames)
        rows: list[ImportedClipRow] = []
        for index, row in enumerate(reader, start=2):
            if not any((value or "").strip() for value in row.values()):
                continue
            rows.append(_row_from_mapping({key.strip().lower(): value for key, value in row.items()}, index))

        return rows


def import_xls_clip_rows(file_path: str | Path) -> list[ImportedClipRow]:
    """Import legacy .xls files when xlrd is available."""

    try:
        import xlrd  # type: ignore[import-not-found]
    except ImportError as error:
        raise ClipImportError(AR_XLS_REQUIRES_CONVERSION) from error

    workbook = xlrd.open_workbook(str(file_path))
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        raise ClipImportError(AR_MISSING_REQUIRED_COLUMNS)

    headers = sheet.row_values(0)
    column_map = _build_column_map(headers)
    imported_rows: list[ImportedClipRow] = []
    for index in range(1, sheet.nrows):
        row = sheet.row_values(index)
        if not _row_has_data(row):
            continue
        imported_rows.append(_row_from_mapping(_values_to_mapping(row, column_map), index + 1))

    return imported_rows


def normalize_imported_time(value: Any) -> str:
    """Normalize spreadsheet time values into HH:MM:SS text."""

    if isinstance(value, datetime):
        return _seconds_to_timestamp(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, time):
        return _seconds_to_timestamp(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, timedelta):
        return _seconds_to_timestamp(round(value.total_seconds()))
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 0 <= value < 1:
            return _seconds_to_timestamp(round(value * 24 * 60 * 60))
        if float(value).is_integer():
            return normalize_timestamp_text(str(int(value)))

    return normalize_timestamp_text(str(value))


def _build_column_map(headers: tuple[Any, ...] | list[Any]) -> dict[str, int]:
    normalized_headers = {str(header).strip().lower(): index for index, header in enumerate(headers) if header is not None}
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in normalized_headers]
    if missing_columns:
        raise ClipImportError(AR_MISSING_REQUIRED_COLUMNS)

    column_map = {column: normalized_headers[column] for column in REQUIRED_COLUMNS}
    for column in OPTIONAL_COLUMNS:
        if column in normalized_headers:
            column_map[column] = normalized_headers[column]

    return column_map


def _values_to_mapping(row: tuple[Any, ...] | list[Any], column_map: dict[str, int]) -> dict[str, Any]:
    return {column: row[index] if index < len(row) else None for column, index in column_map.items()}


def _row_from_mapping(row: dict[str, Any], row_number: int) -> ImportedClipRow:
    try:
        number = _parse_number(row["number"])
        title = "" if row["title"] is None else str(row["title"]).strip()
        start = normalize_imported_time(row["start"])
        end = normalize_imported_time(row["end"])
        exclusions = normalize_imported_exclusions(row.get("exclusions"))
    except (KeyError, TypeError, ValueError) as error:
        raise ClipImportError(f"{AR_INVALID_TIME_VALUE} في الصف {row_number}") from error

    if not title:
        raise ClipImportError(f"العنوان فارغ في الصف {row_number}")

    return ImportedClipRow(number=number, title=title, start=start, end=end, exclusions=exclusions)


def normalize_imported_exclusions(value: Any) -> str:
    """Normalize optional imported exclusion ranges."""

    if value is None or not str(value).strip():
        return ""

    return format_exclusions(parse_exclusions(str(value)))


def _parse_number(value: Any) -> int:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)

    text = normalize_digits(str(value)).strip()
    if not text:
        raise ValueError("Missing number")

    return int(float(text))


def _seconds_to_timestamp(total_seconds: int) -> str:
    if total_seconds < 0:
        raise ValueError("Invalid time")

    return format_seconds(total_seconds)


def _row_has_data(row: tuple[Any, ...] | list[Any]) -> bool:
    return any(value is not None and str(value).strip() for value in row)
