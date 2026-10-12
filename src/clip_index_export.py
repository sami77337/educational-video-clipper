"""Create a complete team handoff package from successfully cut clips.

Only derivative clips and structured indices are included in the package.
No artificial transcripts are generated; transcription stays with the
existing Telegram bot workflow.
"""

from __future__ import annotations

import csv
import json
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.time_utils import format_seconds


INDEX_FILENAME = "FATWA_CLIPS_INDEX.xlsx"
CSV_FILENAME = "FATWA_CLIPS_INDEX.csv"
MANIFEST_FILENAME = "clip_manifest.json"
ZIP_FILENAME = "FATWA_TEAM_PACKAGE.zip"
_HEADERS = (
    "رقم المقطع",
    "عنوان السؤال",
    "البداية الأصلية",
    "النهاية الأصلية",
    "بداية القص الفعلية",
    "نهاية القص الفعلية",
    "نطاق المقطع بالثواني",
    "اسم ملف الفيديو",
    "مسار الملف النسبي",
    "الاستثناءات",
    "مصدر الفيديو",
    "حالة التفريغ",
)


@dataclass(frozen=True)
class HandoffPackage:
    index_path: Path
    csv_path: Path
    manifest_path: Path
    zip_path: Path


def create_handoff_package(
    project_folder: str | Path,
    results: Sequence[Any],
    *,
    source_reference: str = "",
    report_path: str | Path | None = None,
) -> HandoffPackage:
    """Export XLSX/CSV/JSON and a ZIP containing validated MP4 results.

    A missing or unexpected video path is a hard error, never an incomplete
    'success' package. The original input.mp4 is not bundled.
    """
    root = Path(project_folder).resolve()
    if not root.is_dir():
        raise ValueError("مجلد المشروع غير موجود.")
    entries: list[dict[str, Any]] = []
    paths: list[Path] = []
    for item in results:
        clip = item.clip
        output = Path(item.output_path).resolve()
        try:
            relative = output.relative_to(root)
        except ValueError as error:
            raise ValueError("مسار مقطع خارج مجلد المشروع.") from error
        if not output.is_file() or not output.stat().st_size:
            raise ValueError(f"المقطع {clip.number} مفقود أو فارغ؛ لا يمكن إصدار الحزمة.")

        start = getattr(item, "effective_start_seconds", None)
        end = getattr(item, "effective_end_seconds", None)
        if start is None:
            start = clip.start_seconds
        if end is None:
            end = clip.end_seconds

        entries.append({
            "number": clip.number,
            "title": clip.title,
            "source_start": format_seconds(clip.start_seconds),
            "source_end": format_seconds(clip.end_seconds),
            "actual_cut_start": format_seconds(start),
            "actual_cut_end": format_seconds(end),
            "source_span_seconds": round(float(end) - float(start), 3),
            "video_filename": output.name,
            "relative_path": relative.as_posix(),
            "exclusions": clip.exclusions,
            "source_reference": source_reference,
            "transcription_status": "بانتظار التفريغ عبر بوت Telegram",
        })
        paths.append(output)

    if not entries:
        raise ValueError("لا توجد مقاطع ناجحة لإنشاء الحزمة.")

    xlsx_path = root / INDEX_FILENAME
    csv_path = root / CSV_FILENAME
    json_path = root / MANIFEST_FILENAME
    zip_path = root / ZIP_FILENAME

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "فهرس المقاطع"
    sheet.sheet_view.rightToLeft = True
    sheet.append(_HEADERS)
    for entry in entries:
        sheet.append([_safe_cell(entry[key]) for key in (
            "number", "title", "source_start", "source_end",
            "actual_cut_start", "actual_cut_end", "source_span_seconds",
            "video_filename", "relative_path", "exclusions", "source_reference",
            "transcription_status",
        )])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F497D")
        cell.alignment = Alignment(horizontal="center")
    widths = [14, 56, 22, 22, 22, 22, 24, 55, 72, 30, 66, 34]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(xlsx_path)
    workbook.close()

    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(_HEADERS)
        for entry in entries:
            writer.writerow([_safe_cell(entry[key]) for key in (
                "number", "title", "source_start", "source_end",
                "actual_cut_start", "actual_cut_end", "source_span_seconds",
                "video_filename", "relative_path", "exclusions", "source_reference",
                "transcription_status",
            )])
    json_path.write_text(
        json.dumps({
            "schema_version": "1.0",
            "source_reference": source_reference,
            "transcription_provider": "external_telegram_bot",
            "clip_count": len(entries),
            "clips": entries,
        }, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    temp_zip = zip_path.with_suffix(".zip.partial")
    try:
        with ZipFile(temp_zip, "w") as archive:
            report = Path(report_path).resolve() if report_path is not None else None
            extra = [report] if report is not None and report.is_file() and report.is_relative_to(root) else []
            for file_path in (xlsx_path, csv_path, json_path, *extra):
                archive.write(file_path, arcname=file_path.relative_to(root).as_posix(), compress_type=ZIP_DEFLATED)
            for media_path in paths:
                archive.write(
                    media_path,
                    arcname=media_path.relative_to(root).as_posix(),
                    compress_type=ZIP_STORED,
                )
        os.replace(temp_zip, zip_path)
    finally:
        temp_zip.unlink(missing_ok=True)

    return HandoffPackage(
        index_path=xlsx_path, csv_path=csv_path,
        manifest_path=json_path, zip_path=zip_path,
    )


def _safe_cell(value: Any) -> Any:
    """Avoid spreadsheet formula injection from untrusted AI titles."""
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value
