import csv
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from src.clip_index_export import create_handoff_package


def _result(root: Path, title: str = "حكم الصلاة؟", filename: str = "001_حكم الصلاة.mp4"):
    mp4 = root / "مقاطع" / filename
    mp4.parent.mkdir(parents=True, exist_ok=True)
    mp4.write_bytes(b"MP4-SAMPLE" * 250)
    clip = SimpleNamespace(
        number=1, title=title, start_seconds=19, end_seconds=37.5, exclusions=""
    )
    return SimpleNamespace(
        clip=clip,
        output_path=mp4,
        effective_start_seconds=19,
        effective_end_seconds=38,
        destination_folder_name="مقاطع",
    )


def test_team_package_contains_clip_excel_csv_json_and_report(tmp_path: Path):
    result = _result(tmp_path)
    original = tmp_path / "input.mp4"
    original.write_bytes(b"do not package original")
    report = tmp_path / "تقرير القص.txt"
    report.write_text("report", encoding="utf-8")

    package = create_handoff_package(
        tmp_path, [result],
        source_reference="https://www.youtube.com/watch?v=abcDEF12345",
        report_path=report,
    )
    assert package.index_path.is_file()
    assert package.csv_path.is_file()
    data = json.loads(package.manifest_path.read_text(encoding="utf-8"))
    assert data["clip_count"] == 1
    assert data["transcription_provider"] == "external_telegram_bot"
    assert data["clips"][0]["source_end"] == "00:00:37.500"
    assert data["clips"][0]["actual_cut_end"] == "00:00:38"
    assert data["clips"][0]["video_filename"] == result.output_path.name

    workbook = load_workbook(package.index_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        assert sheet.cell(2, 2).value == "حكم الصلاة؟"
        assert sheet.cell(2, 4).value == "00:00:37.500"
        assert sheet.cell(2, 6).value == "00:00:38"
    finally:
        workbook.close()
    with package.csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[1][1] == "حكم الصلاة؟"

    with ZipFile(package.zip_path) as archive:
        entries = archive.namelist()
        assert "FATWA_CLIPS_INDEX.xlsx" in entries
        assert "FATWA_CLIPS_INDEX.csv" in entries
        assert "clip_manifest.json" in entries
        assert "تقرير القص.txt" in entries
        assert "مقاطع/001_حكم الصلاة.mp4" in entries
        assert "input.mp4" not in entries


def test_xlsx_escapes_untrusted_formula_titles(tmp_path: Path):
    result = _result(tmp_path, title='=HYPERLINK("http://example.com")')
    package = create_handoff_package(tmp_path, [result])
    wb = load_workbook(package.index_path, read_only=True)
    try:
        cell = wb.active.cell(2, 2)
        assert cell.data_type != "f"
        assert cell.value.startswith("'=")
    finally:
        wb.close()


def test_handoff_rejects_outside_or_missing_file(tmp_path: Path):
    another = tmp_path.parent / "outside.mp4"
    item = _result(tmp_path)
    item.output_path = another
    with pytest.raises(ValueError, match="خارج"):
        create_handoff_package(tmp_path, [item])
    item.output_path = tmp_path / "not-found.mp4"
    with pytest.raises(ValueError, match="مفقود"):
        create_handoff_package(tmp_path, [item])


def test_empty_clips_rejected(tmp_path: Path):
    with pytest.raises(ValueError, match="لا توجد"):
        create_handoff_package(tmp_path, [])


def test_local_source_path_not_persisted_into_handoff(tmp_path: Path):
    result = _result(tmp_path)
    package = create_handoff_package(tmp_path, [result], source_reference="input.mp4")
    contents = package.manifest_path.read_text(encoding="utf-8")
    assert "input.mp4" in contents
    assert "Users\\SomeName" not in contents



def test_untrusted_title_with_leading_bidi_and_whitespace_is_not_formula(tmp_path: Path):
    result = _result(tmp_path, title="\u200e  =HYPERLINK(\"http://example.com\")")
    package = create_handoff_package(tmp_path, [result])
    workbook = load_workbook(package.index_path, read_only=True, data_only=False)
    try:
        cell = workbook.active.cell(2, 2)
        assert cell.data_type != "f"
        assert cell.value.startswith("'")
    finally:
        workbook.close()
