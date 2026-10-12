"""End-to-end Gemini -> local FFmpeg -> Arabic handoff validation.

Requires real FFmpeg in CI; never substitutes a fake subprocess or fake MP4.
"""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from src.clip_padding import ClipPadding
from src.gemini_import import parse_gemini_response
from src.validation import ClipRowInput
from src.video.ffmpeg_runner import probe_media_duration_seconds, resolve_external_tool
from src.video_processor import VideoProcessor, VideoSourceRequest, VideoSourceType


def _get_ffmpeg() -> str:
    try:
        return resolve_external_tool("ffmpeg")
    except FileNotFoundError:
        if os.environ.get("REQUIRE_FFMPEG_E2E") == "1":
            pytest.fail("FFmpeg must be installed on CI for the real-media integration test.")
        pytest.skip("FFmpeg is not present on this local computer.")


@pytest.fixture
def sample_media(tmp_path: Path) -> Path:
    ffmpeg = _get_ffmpeg()
    output = tmp_path / "source_ar_video.mp4"
    command = [
        ffmpeg, "-hide_banner", "-loglevel", "error", "-y",
        "-f", "lavfi", "-i", "testsrc=size=160x90:rate=25",
        "-f", "lavfi", "-i", "sine=frequency=540:sample_rate=44100",
        "-t", "8", "-c:v", "mpeg4", "-q:v", "5",
        "-pix_fmt", "yuv420p", "-c:a", "aac",
        "-b:a", "64k", str(output),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert completed.returncode == 0, completed.stderr[-2000:]
    assert output.stat().st_size > 1024
    return output


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_local_gemini_video_pipeline_exports_real_mp4_and_handoff(
    tmp_path: Path, sample_media: Path
) -> None:
    ffmpeg = _get_ffmpeg()
    before_hash = _sha256(sample_media)
    gemini = """| م | عنوان السؤال | وقت البداية | وقت النهاية |
| --- | --- | --- | --- |
| 1 | هل يجوز الجمع في السفر؟ | 00:01.250 | 00:03.750 |
| 2 | ما حكم التصدق؟ | 00:04.000 | 00:06.500 |
"""
    parsed = parse_gemini_response(gemini)
    assert len(parsed.clips) == 2
    clips = [
        ClipRowInput(
            row_number=clip.number, title=clip.title,
            start=clip.start, end=clip.end
        )
        for clip in parsed.clips
    ]

    processor = VideoProcessor(output_root=tmp_path / "outputs")
    result = processor.process_project(
        VideoSourceRequest(
            VideoSourceType.LOCAL_FILE,
            str(sample_media),
            source_reference="https://www.youtube.com/watch?v=abcDEF12345",
        ),
        "حزمة تجربة فتوى",
        clips,
        clip_padding=ClipPadding(pre_seconds=0.25, post_seconds=0.25),
        create_handoff_package=True,
    )

    # Verify both outputs are actual decodable media with real durations.
    assert len(result.cut_results) == 2
    for item in result.cut_results:
        assert item.output_path.is_file()
        media_duration = probe_media_duration_seconds(item.output_path)
        assert media_duration == pytest.approx(3.0, abs=0.22)
        decoded = subprocess.run(
            [ffmpeg, "-hide_banner", "-loglevel", "error", "-xerror",
             "-i", str(item.output_path), "-f", "null", "-"],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
        )
        assert decoded.returncode == 0, decoded.stderr[-1000:]

    root = result.project_output_folder
    assert result.export_artifacts.zip_files == [root / "FATWA_TEAM_PACKAGE.zip"]
    assert result.export_artifacts.zip_folder == root
    report_text = result.export_artifacts.report_path.read_text(encoding="utf-8")
    assert "FATWA_TEAM_PACKAGE.zip" in report_text
    manifest = json.loads((root / "clip_manifest.json").read_text(encoding="utf-8"))
    assert manifest["clip_count"] == 2
    assert manifest["source_reference"] == "https://www.youtube.com/watch?v=abcDEF12345"
    assert manifest["clips"][0]["source_reference"] == "https://www.youtube.com/watch?v=abcDEF12345"
    assert manifest["clips"][0]["title"] == "هل يجوز الجمع في السفر؟"
    assert manifest["clips"][0]["source_start"] == "00:00:01.250"
    assert manifest["clips"][0]["source_end"] == "00:00:03.750"
    assert manifest["clips"][0]["actual_cut_start"] == "00:00:01"
    assert manifest["clips"][0]["actual_cut_end"] == "00:00:04"
    assert manifest["clips"][1]["actual_cut_start"] == "00:00:03.750"
    assert manifest["clips"][1]["actual_cut_end"] == "00:00:06.750"

    workbook = load_workbook(root / "FATWA_CLIPS_INDEX.xlsx", read_only=True)
    try:
        sheet = workbook.active
        assert sheet.cell(2, 2).value == "هل يجوز الجمع في السفر؟"
        assert sheet.cell(3, 4).value == "00:00:06.500"
    finally:
        workbook.close()
    with ZipFile(root / "FATWA_TEAM_PACKAGE.zip") as archive:
        contents = archive.namelist()
        assert "FATWA_CLIPS_INDEX.xlsx" in contents
        assert "FATWA_CLIPS_INDEX.csv" in contents
        assert "clip_manifest.json" in contents
        assert len([path for path in contents if path.endswith(".mp4")]) == 2
        assert "input.mp4" not in contents
        assert len([path for path in contents if path.endswith(".txt")]) >= 1
        assert archive.testzip() is None
    assert before_hash == _sha256(sample_media), "The original source must stay unchanged."



def test_gemini_team_package_rejects_times_beyond_source(tmp_path: Path, sample_media: Path):
    processor = VideoProcessor(output_root=tmp_path / "outputs")
    with pytest.raises(Exception, match="يتجاوز مدة الفيديو"):
        processor.process_project(
            VideoSourceRequest(VideoSourceType.LOCAL_FILE, str(sample_media)),
            "خارج حدود المصدر",
            [ClipRowInput(row_number=1, title="سؤال زائد", start="00:07", end="00:09")],
            create_handoff_package=True,
        )
